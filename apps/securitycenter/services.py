import time
import io,re,logging
from pathlib import Path
from django.conf import settings
from django.db import transaction
from django.utils import timezone
from .models import AttachmentScanJob,DLPCase,SecurityIncident
log=logging.getLogger(__name__)

DLP_RULES=[
    ("private_key","critical",re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----",re.I)),
    ("aws_access_key","critical",re.compile(r"\bAKIA[0-9A-Z]{16}\b")),
    ("github_token","critical",re.compile(r"\bgh[pousr]_[A-Za-z0-9_]{30,}\b")),
    ("jwt_token","high",re.compile(r"\beyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\b")),
    ("bearer_token","high",re.compile(r"\bBearer\s+[A-Za-z0-9._~+/=-]{20,}",re.I)),
    ("password_assignment","high",re.compile(r"\b(?:password|passwd|pwd|пароль)\s*[:=]\s*[^\s,;]{6,}",re.I)),
    ("credit_card","high",re.compile(r"(?<!\d)(?:\d[ -]*?){13,19}(?!\d)")),
    ("iban","medium",re.compile(r"\b[A-Z]{2}\d{2}[A-Z0-9]{11,30}\b")),
    ("uzbek_passport","high",re.compile(r"\b[A-Z]{2}\s?\d{7}\b")),
    ("api_secret","high",re.compile(r"\b(?:api[_-]?key|secret[_-]?key|client[_-]?secret)\s*[:=]\s*[A-Za-z0-9_\-./+=]{12,}",re.I)),
    ("confidential_marker","medium",re.compile(r"\b(?:CONFIDENTIAL|STRICTLY CONFIDENTIAL|ДСП|ДЛЯ СЛУЖЕБНОГО ПОЛЬЗОВАНИЯ|КОММЕРЧЕСКАЯ ТАЙНА)\b",re.I)),
]

def _redact(value):
    value=str(value)
    if len(value)<=8:return value[:2]+"***"
    return value[:4]+"…"+value[-4:]

def _luhn_ok(value):
    digits=[int(x) for x in re.sub(r"\D","",value or "")]
    if not 13<=len(digits)<=19:return False
    checksum=0;parity=len(digits)%2
    for i,digit in enumerate(digits):
        if i%2==parity:
            digit*=2
            if digit>9:digit-=9
        checksum+=digit
    return checksum%10==0

def inspect_text(text):
    findings=[]
    for rule,severity,rx in DLP_RULES:
        matches=[]
        for m in rx.finditer(text or ""):
            raw=m.group(0)
            if rule=="credit_card" and not _luhn_ok(raw):continue
            matches.append({"value":_redact(raw),"start":m.start(),"end":m.end()})
            if len(matches)>=10:break
        if matches:findings.append((rule,severity,matches))
    return findings

def _masked_excerpt(text,matches,limit=1200):
    source=(text or "")[:max(limit,1)]
    pieces=[];cursor=0
    for item in sorted(matches,key=lambda x:x["start"]):
        start=max(0,min(int(item["start"]),len(source)));end=max(start,min(int(item["end"]),len(source)))
        if start<cursor:continue
        pieces.append(source[cursor:start]);pieces.append(item["value"]);cursor=end
    pieces.append(source[cursor:])
    return "".join(pieces)[:limit]

def scan_message_dlp(message):
    for rule,severity,matches in inspect_text(message.body):
        if DLPCase.objects.filter(message=message,rule=rule).exists():continue
        case=DLPCase.objects.create(
            source=DLPCase.Source.MESSAGE,severity=severity,sender=message.sender,
            conversation=message.conversation,message=message,rule=rule,
            matches=matches,excerpt=_masked_excerpt(message.body,matches),
        )
        SecurityIncident.objects.create(kind=SecurityIncident.Kind.DLP,user=message.sender,severity=severity,title=f"DLP: {rule}",details={"case_id":case.pk,"message_id":message.pk})

def queue_attachment_security_scan(attachment,force=False):
    job,created=AttachmentScanJob.objects.get_or_create(attachment=attachment)
    if force or job.status==AttachmentScanJob.Status.ERROR:
        job.status=AttachmentScanJob.Status.QUEUED;job.attempts=0;job.last_error="";job.started_at=None;job.finished_at=None
        job.save(update_fields=["status","attempts","last_error","started_at","finished_at","updated_at"])
    return job

def _clamd_client():
    import clamd
    return clamd.ClamdNetworkSocket(host=getattr(settings,"CLAMAV_HOST","clamav"),port=int(getattr(settings,"CLAMAV_PORT",3310)),timeout=60)

def _scan_clamav(attachment):
    import os,tempfile
    suffix=Path(attachment.original_name or getattr(attachment.file,"name","")).suffix[:16] or ".bin"
    tmp=tempfile.NamedTemporaryFile(delete=False,suffix=suffix)
    try:
        attachment.file.open("rb")
        try:
            for chunk in iter(lambda:attachment.file.read(1024*1024),b""):
                tmp.write(chunk)
        finally:
            attachment.file.close();tmp.flush();tmp.close()
        with open(tmp.name,"rb") as fh:
            result=_clamd_client().instream(fh)
    finally:
        try:tmp.close()
        except Exception:pass
        try:os.unlink(tmp.name)
        except OSError:pass
    status,signature=(result or {}).get("stream",("ERROR","no response"))
    attachment.scan_engine="ClamAV";attachment.scanned_at=timezone.now()
    if status=="OK":
        attachment.scan_status="safe";attachment.scan_signature=""
    elif status=="FOUND":
        attachment.scan_status="infected";attachment.scan_signature=str(signature)[:255]
        SecurityIncident.objects.create(kind=SecurityIncident.Kind.MALWARE,user=attachment.message.sender,severity="critical",title=f"Malware detected: {signature}",details={"attachment_id":attachment.pk,"message_id":attachment.message_id,"sha256":attachment.sha256})
    else:
        raise RuntimeError(f"ClamAV returned {status}: {signature}")
    attachment.save(update_fields=["scan_status","scan_engine","scan_signature","scanned_at"])

def _extract_document_text(attachment):
    name=(attachment.original_name or "").lower()
    ct=(attachment.content_type or "").lower()
    max_bytes=int(getattr(settings,"DLP_MAX_PARSE_MB",25))*1024*1024
    if attachment.size and attachment.size>max_bytes:return ""
    attachment.file.open("rb")
    try:data=attachment.file.read(max_bytes+1)
    finally:attachment.file.close()
    if len(data)>max_bytes:return ""

    if ct.startswith("text/") or any(name.endswith(x) for x in (".txt",".csv",".json",".xml",".yaml",".yml",".ini",".conf",".env",".log",".py",".js",".ts",".ps1",".sh")):
        return data.decode("utf-8",errors="ignore")[:200000]

    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages[:100])[:200000]
        if name.endswith(".docx"):
            from docx import Document
            doc=Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)[:200000]
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            lines=[];size=0
            for ws in wb.worksheets[:20]:
                for row in ws.iter_rows(values_only=True):
                    line=" | ".join("" if v is None else str(v) for v in row)
                    if line.strip():
                        lines.append(line);size+=len(line)
                    if size>200000:break
                if size>200000:break
            return "\n".join(lines)[:200000]
        if name.endswith(".pptx"):
            from pptx import Presentation
            deck=Presentation(io.BytesIO(data));lines=[]
            for slide in deck.slides[:100]:
                for shape in slide.shapes:
                    if hasattr(shape,"text") and shape.text:lines.append(shape.text)
            return "\n".join(lines)[:200000]
    except Exception:
        log.exception("DLP document text extraction failed for attachment %s",attachment.pk)
    return ""

def _scan_attachment_dlp(attachment):
    text=_extract_document_text(attachment)
    if not text:return
    for rule,severity,matches in inspect_text(text):
        if DLPCase.objects.filter(attachment=attachment,rule=rule).exists():continue
        case=DLPCase.objects.create(
            source=DLPCase.Source.ATTACHMENT,severity=severity,sender=attachment.message.sender,
            conversation=attachment.message.conversation,message=attachment.message,
            attachment=attachment,rule=rule,matches=matches,excerpt=_masked_excerpt(text,matches),
        )
        SecurityIncident.objects.create(kind=SecurityIncident.Kind.DLP,user=attachment.message.sender,severity=severity,title=f"DLP file: {rule}",details={"case_id":case.pk,"attachment_id":attachment.pk})

def recover_stale_jobs(stale_seconds=120):
    from datetime import timedelta
    from apps.chat.models import Attachment
    cutoff=timezone.now()-timedelta(seconds=max(30,int(stale_seconds)))
    reset=0;created=0
    for job in AttachmentScanJob.objects.filter(status=AttachmentScanJob.Status.PROCESSING,updated_at__lt=cutoff).select_related("attachment"):
        job.status=AttachmentScanJob.Status.QUEUED;job.started_at=None;job.finished_at=None;job.last_error="stale job automatically requeued"
        job.save(update_fields=["status","started_at","finished_at","last_error","updated_at"]);reset+=1
    for a in Attachment.objects.filter(scan_status__in=["pending","error"]).iterator(chunk_size=200):
        job=AttachmentScanJob.objects.filter(attachment=a).first()
        if not job:
            queue_attachment_security_scan(a,force=True);created+=1
        elif job.status==AttachmentScanJob.Status.ERROR and job.updated_at<cutoff:
            queue_attachment_security_scan(a,force=True);reset+=1
    return reset,created

def claim_next_job():
    with transaction.atomic():
        job=(AttachmentScanJob.objects.select_for_update(skip_locked=True).select_related("attachment__message__sender","attachment__message__conversation").filter(status=AttachmentScanJob.Status.QUEUED).order_by("queued_at").first())
        if not job:return None
        job.status=AttachmentScanJob.Status.PROCESSING;job.started_at=timezone.now();job.attempts+=1;job.save(update_fields=["status","started_at","attempts","updated_at"]);return job

def _broadcast_attachment_state(attachment):
    try:
        from apps.chat.services import broadcast,message_payload,broadcast_message_update_to_user_clients
        msg=attachment.message
        payload=message_payload(msg)
        broadcast(msg.conversation_id,"message_updated",payload)
        broadcast_message_update_to_user_clients(msg)
    except Exception:
        log.exception("Unable to broadcast security scan result")

def run_one_job():
    job=claim_next_job()
    if not job:return False
    a=job.attachment
    try:
        _scan_clamav(a)
        if a.scan_status=="safe":_scan_attachment_dlp(a)
        job.status=AttachmentScanJob.Status.DONE;job.finished_at=timezone.now();job.last_error="";job.save(update_fields=["status","finished_at","last_error","updated_at"])
        _broadcast_attachment_state(a)
    except Exception as exc:
        log.exception("Security scan failed for attachment %s",a.pk)
        job.last_error=str(exc)[:4000];job.finished_at=timezone.now()
        if job.attempts<12:
            # ClamAV can take tens of seconds to warm up after a host reboot. Keep the
            # attachment pending and retry instead of permanently stranding it after 3 fast attempts.
            job.status=AttachmentScanJob.Status.QUEUED
            a.scan_status="pending";a.scan_engine="ClamAV";a.scanned_at=timezone.now()
            a.save(update_fields=["scan_status","scan_engine","scanned_at"])
            job.save(update_fields=["status","finished_at","last_error","updated_at"])
            time.sleep(min(8.0,1.25*max(1,job.attempts)))
        else:
            job.status=AttachmentScanJob.Status.ERROR
            a.scan_status="error";a.scan_engine="ClamAV";a.scanned_at=timezone.now()
            a.save(update_fields=["scan_status","scan_engine","scanned_at"])
            job.save(update_fields=["status","finished_at","last_error","updated_at"])
            _broadcast_attachment_state(a)
    return True
