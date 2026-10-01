import io
from pathlib import Path
from django.conf import settings
from django.urls import reverse


def build_preview(attachment):
    if attachment.scan_status=="infected":
        return {"kind":"blocked","title":attachment.original_name,"text":"Файл заблокирован проверкой безопасности."}
    name=(attachment.original_name or "file").lower();ct=(attachment.content_type or "").lower()
    max_bytes=int(getattr(settings,"FILE_PREVIEW_MAX_MB",25))*1024*1024
    if attachment.size and attachment.size>max_bytes:return {"kind":"large","title":attachment.original_name,"text":f"Preview ограничен {settings.FILE_PREVIEW_MAX_MB} MB."}
    content_url=reverse("chat:attachment_content",args=[attachment.pk])
    if attachment.is_image:return {"kind":"image","title":attachment.original_name,"url":content_url}
    if attachment.is_video:return {"kind":"video","title":attachment.original_name,"url":content_url}
    if attachment.is_audio:return {"kind":"audio","title":attachment.original_name,"url":content_url}
    attachment.file.open("rb")
    try:data=attachment.file.read(max_bytes+1)
    finally:attachment.file.close()
    if len(data)>max_bytes:return {"kind":"large","title":attachment.original_name,"text":"Файл слишком большой для inline preview."}
    try:
        if ct.startswith("text/") or Path(name).suffix in {".txt",".csv",".json",".xml",".yaml",".yml",".ini",".conf",".env",".log",".py",".js",".ts",".ps1",".sh"}:
            return {"kind":"text","title":attachment.original_name,"text":data.decode("utf-8",errors="replace")[:120000]}
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(data));pages=[]
            for idx,page in enumerate(reader.pages[:40],1):pages.append({"number":idx,"text":(page.extract_text() or "")[:15000]})
            return {"kind":"pdf","title":attachment.original_name,"pages":pages}
        if name.endswith(".docx"):
            from docx import Document
            doc=Document(io.BytesIO(data));text="\n".join(p.text for p in doc.paragraphs)
            return {"kind":"document","title":attachment.original_name,"text":text[:120000]}
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True);sheets=[]
            for ws in wb.worksheets[:12]:
                rows=[]
                for row in ws.iter_rows(values_only=True):
                    rows.append(["" if x is None else str(x)[:300] for x in row[:30]])
                    if len(rows)>=120:break
                sheets.append({"name":ws.title,"rows":rows})
            return {"kind":"spreadsheet","title":attachment.original_name,"sheets":sheets}
        if name.endswith(".pptx"):
            from pptx import Presentation
            deck=Presentation(io.BytesIO(data));slides=[]
            for i,slide in enumerate(deck.slides[:60],1):
                text=[]
                for shape in slide.shapes:
                    if hasattr(shape,"text") and shape.text:text.append(shape.text)
                slides.append({"number":i,"text":"\n".join(text)[:12000]})
            return {"kind":"presentation","title":attachment.original_name,"slides":slides}
    except Exception as exc:
        return {"kind":"error","title":attachment.original_name,"text":f"Preview не удалось построить: {exc}"}
    return {"kind":"download","title":attachment.original_name,"url":f"{content_url}?download=1","text":"Файл готов к скачиванию."}
